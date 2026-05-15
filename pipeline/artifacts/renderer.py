"""Blueprint Renderer — turn a ReferenceAnswer into a real deliverable file.

This replaces the old best-of-n "model solves task" approach with deterministic
code rendering. The blueprint was designed synchronously with the task, so the
prompt, rubric, and answer are all consistent.

For formats that need narrative text (docx, md), the paragraphs are pre-generated
during synthesis and stored in the blueprint. Rendering is pure code assembly.
"""

from __future__ import annotations

import logging
import shutil
import subprocess
from pathlib import Path

from pipeline.artifacts.base import deliverable_dir
from pipeline.config import DATA_DIR
from pipeline.scenario.reference_answer import ReferenceAnswer
from pipeline.scenario.task_candidate import TaskCandidate

logger = logging.getLogger(__name__)

GOLD_DIR = DATA_DIR / "gold"
GOLD_DIR.mkdir(parents=True, exist_ok=True)


def maybe_render_pdf(docx_path: Path) -> Path | None:
    """Use LibreOffice headless to convert .docx → .pdf if available."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        logger.info("LibreOffice not found; skipping PDF rendering")
        return None
    try:
        subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to", "pdf",
                "--outdir", str(docx_path.parent),
                str(docx_path),
            ],
            check=True,
            capture_output=True,
            timeout=60,
        )
        pdf_path = docx_path.with_suffix(".pdf")
        if pdf_path.exists():
            logger.info("  PDF rendered → %s", pdf_path)
            return pdf_path
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        logger.warning("PDF rendering failed: %s", e)
    return None


# ─────────────────────────── XLSX Rendering ───────────────────────────


def _render_xlsx(blueprint: ReferenceAnswer, out_path: Path) -> Path:
    """Render an xlsx from a SheetBlueprint."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    FILL_COLORS = {
        "header": "1F4E78",
        "input": "FFF2CC",
        "subtotal": "D9E1F2",
        "total": "BDD7EE",
    }

    def apply_style(cell, spec):
        if spec.bold or spec.italic:
            cell.font = Font(
                name="Calibri", size=11,
                bold=spec.bold, italic=spec.italic,
                color="FFFFFF" if spec.fill == "header" else "000000",
            )
        if spec.fill != "none":
            color = FILL_COLORS.get(spec.fill)
            if color:
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        if spec.number_format:
            cell.number_format = spec.number_format
        if spec.alignment:
            cell.alignment = Alignment(horizontal=spec.alignment, vertical="center")

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_plan in (blueprint.sheets or []):
        name = sheet_plan.name.replace("/", "-").replace("\\", "-")[:31]
        ws = wb.create_sheet(title=name)

        for spec in sheet_plan.cells:
            try:
                cell = ws[spec.ref]
            except (ValueError, AttributeError):
                logger.warning("invalid cell ref %s; skipping", spec.ref)
                continue
            if spec.formula is not None:
                cell.value = spec.formula if spec.formula.startswith("=") else f"={spec.formula}"
            elif spec.value is not None:
                cell.value = spec.value
            apply_style(cell, spec)

        for col_letter, width in sheet_plan.column_widths.items():
            try:
                ws.column_dimensions[col_letter.upper()].width = max(8.0, min(60.0, width))
            except (ValueError, KeyError):
                pass

        if not sheet_plan.column_widths:
            ws.column_dimensions["A"].width = 40
            for c in range(2, 10):
                ws.column_dimensions[get_column_letter(c)].width = 16

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    logger.info("wrote .xlsx → %s", out_path)
    return out_path


# ─────────────────────────── DOCX Rendering ───────────────────────────


def _add_para_with_bold(doc, text: str, *, alignment=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    parts = text.split("**")
    bold = False
    for part in parts:
        if part:
            run = p.add_run(part)
            run.bold = bold
        bold = not bold
    return p


def _render_docx(blueprint: ReferenceAnswer, out_path: Path) -> Path:
    """Render a docx from a SectionBlueprint."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    if blueprint.page_settings == "double_spaced":
        for sect in doc.sections:
            sect.top_margin = sect.bottom_margin = sect.left_margin = sect.right_margin = Pt(72)
        style.paragraph_format.line_spacing = 2.0

    if blueprint.title:
        h = doc.add_heading(blueprint.title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if blueprint.doc_header:
        hdr = blueprint.doc_header
        for label, value in [
            ("TO", hdr.to), ("FROM", hdr.from_), ("DATE", hdr.date),
            ("RE", hdr.re), ("CC", hdr.cc),
        ]:
            if value:
                p = doc.add_paragraph()
                run = p.add_run(f"{label}:\t")
                run.bold = True
                p.add_run(value)
        doc.add_paragraph("_" * 80)

    for section in (blueprint.sections or []):
        if section.heading:
            doc.add_heading(section.heading, level=section.heading_level)
        for para in section.paragraphs:
            _add_para_with_bold(doc, para)

    if blueprint.signature_block:
        doc.add_paragraph()
        for line in blueprint.signature_block.split("\n"):
            doc.add_paragraph(line)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    logger.info("wrote .docx → %s", out_path)
    return out_path


# ─────────────────────────── Markdown Rendering ───────────────────────────


def _render_markdown(blueprint: ReferenceAnswer, out_path: Path) -> Path:
    """Render a markdown document from MdSectionBlueprint."""
    lines = []
    if blueprint.md_title:
        lines.append(f"# {blueprint.md_title}")
        lines.append("")

    for section in (blueprint.md_sections or []):
        # Skip sections whose heading duplicates the md_title
        if blueprint.md_title and section.heading.strip() == blueprint.md_title.strip():
            continue
        # Clamp heading_level to min 2 when md_title already rendered as H1
        level = section.heading_level
        if blueprint.md_title and level < 2:
            level = 2
        prefix = "#" * level
        lines.append(f"{prefix} {section.heading}")
        lines.append("")
        for para in section.paragraphs:
            lines.append(para)
            lines.append("")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(lines), encoding="utf-8")
    logger.info("wrote .md → %s", out_path)
    return out_path


# ─────────────────────────── Public API ───────────────────────────


def render_deliverable(task: TaskCandidate) -> Path:
    """Render the reference answer blueprint to a real file."""
    blueprint = task.canonical.reference_answer
    if not blueprint:
        raise ValueError(f"No reference_answer for {task.candidate_id}")

    out_dir = deliverable_dir(task.candidate_id)
    fmt = blueprint.format

    if fmt == "xlsx":
        out_path = out_dir / f"{task.deliverable.deliverable_id}.xlsx"
        return _render_xlsx(blueprint, out_path)
    elif fmt == "docx":
        out_path = out_dir / f"{task.deliverable.deliverable_id}.docx"
        return _render_docx(blueprint, out_path)
    elif fmt == "pdf":
        docx_path = out_dir / f"{task.deliverable.deliverable_id}.docx"
        pdf_path = out_dir / f"{task.deliverable.deliverable_id}.pdf"
        _render_docx(blueprint, docx_path)
        converted = maybe_render_pdf(docx_path)
        if converted:
            return converted
        return docx_path  # fallback if LibreOffice unavailable
    elif fmt in ("markdown", "md"):
        out_path = out_dir / f"{task.deliverable.deliverable_id}.md"
        return _render_markdown(blueprint, out_path)
    else:
        # Fallback to markdown for unsupported formats
        out_path = out_dir / f"{task.deliverable.deliverable_id}.md"
        return _render_markdown(blueprint, out_path)


def materialize_deliverable(task: TaskCandidate) -> dict:
    """Render and persist the deliverable. Returns metadata dict."""
    out_path = render_deliverable(task)
    return {
        "candidate_id": task.candidate_id,
        "format": task.canonical.reference_answer.format if task.canonical.reference_answer else "unknown",
        "file_path": str(out_path),
        "file_size": out_path.stat().st_size if out_path.exists() else 0,
    }

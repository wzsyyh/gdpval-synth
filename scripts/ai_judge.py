#!/usr/bin/env python3
"""AI Judge: automatically evaluate deliverables against their rubrics
by reading the deliverable content and checking each rubric item.

Usage: uv run python scripts/ai_judge.py [--task TASK_ID]
"""

import sys
import json
import argparse
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

ACCEPTED_DIR = Path("data/accepted")
DELIVERABLES_DIR = Path("data/deliverables")

# ── Read raw content from deliverables ──

def read_docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for i, t in enumerate(doc.tables):
        parts.append(f"\n[TABLE {i}]")
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def read_xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=False)
    parts = []
    for name in wb.sheetnames:
        ws = wb[name]
        parts.append(f"\n=== Sheet: {name} ===")
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 100)):
            cells = []
            for c in row:
                if c.value is not None:
                    v = str(c.value)
                    if len(v) > 80:
                        v = v[:77] + "..."
                    cells.append(f"{c.coordinate}={v}")
            if cells:
                parts.append("  " + " | ".join(cells))
    return "\n".join(parts)


def read_md_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def get_deliverable_text(path: Path) -> str:
    fmt = path.suffix.lower()
    try:
        if fmt == ".docx":
            return read_docx_text(path)
        elif fmt == ".xlsx":
            return read_xlsx_text(path)
        elif fmt == ".md":
            return read_md_text(path)
        else:
            return f"[Unsupported format: {fmt}]"
    except Exception as e:
        return f"[Error reading {path.name}]: {e}"


# ── Keyword-based matching helpers ──

def text_contains(text: str, keywords: list[str], *, all_words: bool = False) -> bool:
    """Check if text contains any (or all) of the keywords."""
    text_lower = text.lower()
    if all_words:
        return all(kw.lower() in text_lower for kw in keywords)
    return any(kw.lower() in text_lower for kw in keywords)


def text_contains_numbers(text: str, n: int = 1) -> bool:
    """Check if text contains at least n numbers."""
    import re
    return len(re.findall(r'\b\d+(?:[.,]\d+)?(?:%|\s*(?:billion|million|M|B))?', text)) >= n


def get_sheets_names(xlsx_path: Path) -> list[str]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(str(xlsx_path), data_only=False)
        return wb.sheetnames
    except Exception:
        return []


def count_formulas(xlsx_path: Path) -> int:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(str(xlsx_path), data_only=False)
        count = 0
        for name in wb.sheetnames:
            ws = wb[name]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        count += 1
        return count
    except Exception:
        return 0


# ── Occupation-specific checker hints ──

OCCUPATION_HINTS = {
    "lawyer": {
        "section_headers": ["Introduction", "Background", "Analysis", "Argument", "Conclusion"],
        "expected_terms": ["court", "plaintiff", "defendant", "jurisdiction", "precedent", "pursuant"],
    },
    "financial_analyst": {
        "section_headers": ["Executive Summary", "Overview", "Analysis", "Recommendation"],
        "expected_terms": ["revenue", "EBITDA", "margin", "cash flow", "debt", "$"],
    },
    "software_engineer": {
        "section_headers": ["Executive Summary", "Overview", "Assessment", "Recommendation"],
        "expected_terms": ["PR", "commit", "code review", "test", "performance", "security"],
    },
}


def judge_item(criterion: str, cat: str, text: str, deliverable_path: Path,
               task: dict, summary: dict) -> tuple[str, str]:
    """Judge a single rubric item against deliverable content.
    Returns (judgment: ✓/✗/△, reason).
    """
    text_lower = text.lower()
    criterion_lower = criterion.lower()

    # ── FORMAT items ──
    if cat == "format":
        if "docx" in criterion_lower or "word" in criterion_lower:
            if deliverable_path.suffix == ".docx":
                return ("✓", "File is .docx")
            return ("✗", f"File is {deliverable_path.suffix}, not docx")
        if "xlsx" in criterion_lower or "excel" in criterion_lower or "workbook" in criterion_lower:
            if deliverable_path.suffix == ".xlsx":
                return ("✓", "File is .xlsx")
            return ("✗", f"File is {deliverable_path.suffix}, not xlsx")
        if "markdown" in criterion_lower or ".md" in criterion_lower:
            if deliverable_path.suffix == ".md":
                return ("✓", "File is .md")
            return ("✗", f"File is {deliverable_path.suffix}, not md")
        if "track changes" in criterion_lower or "revision" in criterion_lower:
            return ("△", "Track changes state not detectable from text extraction (need to open in Word)")
        if "file name" in criterion_lower or "naming" in criterion_lower or "protocol" in criterion_lower:
            # File naming check
            fname = deliverable_path.name
            if any(kw in fname for kw in ["motion", "memo", "outline", "redline", "review", "credit", "model", "analysis"]):
                return ("△", f"File name '{fname}' seems reasonable, manual check needed for exact protocol")
            return ("✗", f"File name '{fname}' doesn't match expected convention")
        if "one page" in criterion_lower or "page or less" in criterion_lower:
            # Rough check by character count
            if len(text) < 3000:
                return ("✓", "Content is short enough for one page")
            return ("△", f"Content is {len(text):,} chars, might be more than one page")

    # ── STRUCTURE items ──
    if cat == "structure":
        if "section" in criterion_lower or "heading" in criterion_lower or "template" in criterion_lower:
            # Check for required section
            for kw in ["executive summary", "introduction", "background", "analysis",
                       "argument", "conclusion", "recommendation", "overview",
                       "risk factor", "financial analysis", "code review", "assessment"]:
                if kw in criterion_lower:
                    # Check if section heading exists
                    import re
                    headings = re.findall(r'^(.+)$', text, re.MULTILINE)
                    for h in headings:
                        if kw.lower() in h.lower() or kw.lower().rstrip("s") in h.lower():
                            return ("✓", f"Section '{kw}' found: '{h.strip()}'")
                    # Check in text body too
                    if kw in text_lower:
                        return ("△", f"'{kw}' mentioned in text but not as clear section heading")
                    return ("✗", f"Section '{kw}' not found")
            # If no specific keyword match, check for general section-like structure
            import re
            headings = re.findall(r'^.{0,3}(?:[A-Z][a-z]+(?:\s+[A-Z][a-z]+)+)', text, re.MULTILINE)
            if len(headings) >= 3:
                return ("✓", f"Found {len(headings)} potential section headings")

    # ── REFERENCE items ──
    if cat == "reference":
        if "cite" in criterion_lower or "citation" in criterion_lower or "case" in criterion_lower or "precedent" in criterion_lower:
            # Look for legal citation patterns
            import re
            citations = re.findall(r'\d+\s+[A-Za-z\.]+\s+\d+', text)
            us_citations = re.findall(r'\d+\s+U\.S\.\s+\d+', text)
            f_citations = re.findall(r'\d+\s+F\.\d[dh]?\s+\d+', text)
            all_cites = citations + us_citations + f_citations
            if all_cites:
                return ("✓", f"Found citations: {all_cites[:3]}")
            # Check specific named references
            if "v." in text_lower or "v " in text_lower:
                return ("△", "Mentions case names but no standard citation format detected")
            return ("✗", "No legal citations found")
        if "section" in criterion_lower or "statute" in criterion_lower or "regulation" in criterion_lower:
            # Check for statute references
            import re
            statutes = re.findall(r'\d+\s+U\.S\.C\.', text)
            if statutes:
                return ("✓", f"Found statute references: {statutes}")
        if any(kw in criterion_lower for kw in ["identifies", "names", "lists", "references", "refers to"]):
            # Extract what should be identified
            import re
            # Try to find nouns that look like entities
            entities = ["amazon", "meridian", "saltspring", "plaintiff", "defendant", "amzn", "cardinal"]
            for ent in entities:
                if ent in text_lower:
                    return ("✓", f"Contains '{ent}'")
            return ("△", "Check exact entity match manually")

    # ── ACCURACY / DATA items ──
    if cat in ("accuracy", "data"):
        import re
        # Check for specific numbers
        numbers = re.findall(r'\$?\d+(?:[,.]\d+)?\s*(?:million|billion|M|B|%|basis\s*points|bps)?', text_lower)
        if numbers:
            return ("△", f"Found numbers: {numbers[:5]} — exact value match needs manual verification")
        # Check for named entities
        for kw in ["amazon", "amzn", "meridian", "saltspring", "cardinal", "kimble", "paxton"]:
            if kw in criterion_lower and kw in text_lower:
                return ("✓", f"Entity '{kw}' found in text")
        if "correctly" in criterion_lower or "accurate" in criterion_lower:
            return ("△", "Needs manual verification for correctness")

    # ── CONTENT items ──
    if cat == "content":
        # Extract meaningful keywords from criterion
        import re
        # Skip very generic criteria
        generic = ["professional", "standard", "appropriate", "proper", "clear", "comprehensive",
                   "all required", "necessary"]
        if any(g in criterion_lower for g in generic) and len(criterion) < 40:
            return ("△", "Generic criterion, needs manual judgment")

        # Look for content keywords in the text
        words = re.findall(r'[A-Z][a-z]+(?:\s+[A-Z][a-z]+)*', criterion)
        specific_words = [w for w in words if len(w) > 4 and w.lower() not in
                         ["Background", "Introduction", "Conclusion", "Executive", "Summary",
                          "Analysis", "Section", "Content", "Requirement"]]
        if specific_words:
            found = [w for w in specific_words if w.lower() in text_lower]
            if found:
                return ("✓", f"Content keywords found: {found[:3]}")
            if len(specific_words) <= 3:
                return ("✗", f"Expected content not found: {specific_words}")
            return ("△", f"Could not verify all content keywords: {specific_words}")

    # ── STYLE items ──
    if cat == "style":
        if "professional" in criterion_lower:
            # Check for AI tells
            ai_tells = ["navigate the complex", "in today's", "delve", "ensure that",
                       "it is important", "landscape", "robust", "cutting-edge"]
            found_tells = [t for t in ai_tells if t in text_lower]
            if found_tells:
                return ("△", f"Contains potential AI tells: {found_tells}")
            return ("✓", "No obvious AI tells, text reads professionally")
        if "formatting" in criterion_lower or "consistent" in criterion_lower:
            return ("△", "Formatting consistency needs manual review")
        if "defined term" in criterion_lower or "capitalized" in criterion_lower:
            import re
            defined = re.findall(r'[A-Z][a-z]+(?:\s+[A-Z][a-z]+)*\s+[Mm]eans?', text[:3000])
            if defined:
                return ("✓", f"Found defined terms: {defined[:3]}")
            return ("△", "Check capitalization conventions manually")

    # ── COMPLETENESS items ──
    if cat == "completeness":
        if "summariz" in criterion_lower or "include" in criterion_lower or "cover" in criterion_lower:
            # Extract what should be included
            import re
            items = re.findall(r'(?:summariz|includ|cover)(?:es|ed|ing)?\s+(?:the\s+)?(?:addition|removal|revision|change|update|modification|section|clause|provision)', criterion_lower)
            if items:
                # Check if this content is in the text
                for kw in ["addition", "removal", "revision", "change", "amendment",
                          "confidential", "indemnif", "termination", "representation",
                          "compliance", "carve-out", "waiver", "limitation"]:
                    if kw in criterion_lower and kw in text_lower:
                        return ("✓", f"'{kw}' mentioned in text")
                return ("△", "Check if this summary content is present manually")
            if len(text) > 1000:
                return ("✓", "Document has substantial content")
            return ("✗", "Document content is too sparse")

    # ── JUDGMENT items ──
    if cat == "judgment":
        return ("△", "Subjective judgment item, needs human review")

    # Fallback
    return ("△", "Could not auto-judge, needs manual review")


def judge_deliverable(task_id: str, task: dict, deliverable_path: Path):
    rubric = task.get("rubric", [])
    prompt = task.get("prompt", "")
    occupation = task.get("occupation", "?")
    archetype = task.get("archetype", "?")
    fmt = deliverable_path.suffix.lstrip(".")

    text = get_deliverable_text(deliverable_path)

    total_pts = sum(r["score"] for r in rubric)
    max_possible = sum(r["score"] for r in rubric if r["score"] > 0)

    # Judge each item
    results = []
    earned = 0
    max_score = 0
    auto_judged = 0
    needs_manual = 0

    for r in rubric:
        score = r["score"]
        cat = r.get("category", "content")
        criterion = r.get("criterion", "")
        is_penalty = r.get("is_penalty", False)

        if score > 0:
            max_score += score

        judgment, reason = judge_item(criterion, cat, text, deliverable_path, task, {})

        # Assign score
        if judgment == "✓":
            earned += score
            auto_judged += 1
        elif judgment == "△":
            needs_manual += 1
            # For "△" we conservatively don't award points

        results.append({
            "score": score,
            "criterion": criterion,
            "category": cat,
            "is_penalty": is_penalty,
            "judgment": judgment,
            "reason": reason,
        })

    # Summary by category
    cat_stats = {}
    for r in results:
        cat = r["category"]
        if cat not in cat_stats:
            cat_stats[cat] = {"total": 0, "pass": 0, "partial": 0, "fail": 0, "unknown": 0}
        cat_stats[cat]["total"] += 1
        if r["judgment"] == "✓":
            cat_stats[cat]["pass"] += 1
        elif r["judgment"] == "△":
            cat_stats[cat]["partial"] += 1
        elif r["judgment"] == "✗":
            cat_stats[cat]["fail"] += 1
        else:
            cat_stats[cat]["unknown"] += 1

    print(f"\n{'='*80}")
    print(f"  JUDGE REPORT: {task_id}")
    print(f"  {occupation} — {archetype} ({fmt.upper()})")
    print(f"  File: {deliverable_path.name} ({deliverable_path.stat().st_size:,} bytes)")
    print(f"{'='*80}")

    print(f"\n  Auto-judged: {auto_judged}/{len(rubric)} items | Needs manual: {needs_manual}")
    print(f"  Auto-earned: {earned}/{max_score} points (from auto-judged items)")

    # Show results grouped by verdict
    for verdict in ["✓", "△", "✗"]:
        items = [r for r in results if r["judgment"] == verdict]
        if not items:
            continue
        label = {"✓": "PASS", "△": "NEEDS REVIEW", "✗": "FAIL"}[verdict]
        print(f"\n  ── {label} ({len(items)} items) ──")
        for r in items:
            s = r["score"]
            score_str = f"+{s}" if s > 0 else str(s)
            cat = r["category"]
            reason = r["reason"]
            print(f"    [{score_str}] ({cat}) {r['criterion'][:80]}")
            print(f"           → {reason}")

    print(f"\n  ── Summary by Category ──")
    print(f"  {'Category':<16} {'Total':>5} {'Pass':>5} {'Partial':>7} {'Fail':>5}")
    print(f"  {'-'*42}")
    for cat, st in sorted(cat_stats.items(), key=lambda x: -x[1]["total"]):
        print(f"  {cat:<16} {st['total']:>5} {st['pass']:>5} {st['partial']:>7} {st['fail']:>5}")

    # Quick stats
    n_pass = sum(1 for r in results if r["judgment"] == "✓")
    n_partial = sum(1 for r in results if r["judgment"] == "△")
    n_fail = sum(1 for r in results if r["judgment"] == "✗")

    print(f"\n  ── Overall ──")
    print(f"  ✓ {n_pass} passed | △ {n_partial} needs review | ✗ {n_fail} failed")
    print(f"  Auto-earned: {earned}/{max_score} pts (needs manual review for {needs_manual} items)")

    return results


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--task", type=str)
    parser.add_argument("--occupation", type=str)
    args = parser.parse_args()

    task_files = sorted(ACCEPTED_DIR.glob("sc_*.json"))
    if not task_files:
        print("No accepted tasks found.")
        return

    if args.task:
        task_files = [tf for tf in task_files if tf.stem == args.task]
    if args.occupation:
        task_files = [tf for tf in task_files
                      if json.loads(tf.read_text()).get("occupation") == args.occupation]

    all_results = {}
    for tf in task_files:
        task_id = tf.stem
        task = json.loads(tf.read_text(encoding="utf-8"))

        deliv_dir = DELIVERABLES_DIR / task_id
        if not deliv_dir.exists() or not list(deliv_dir.iterdir()):
            print(f"\n⚠  {task_id}: No deliverable — SKIPPED")
            continue

        for f in list(deliv_dir.iterdir()):
            results = judge_deliverable(task_id, task, f)
            all_results[f"{task_id}/{f.name}"] = results

    # Final overview
    print(f"\n{'='*80}")
    print(f"  FINAL OVERVIEW")
    print(f"{'='*80}")
    print(f"  {'File':<50} {'Items':>5} {'Auto✓':>5} {'△':>3} {'✗':>3}")
    print(f"  {'-'*70}")
    for name, results in sorted(all_results.items()):
        n_pass = sum(1 for r in results if r["judgment"] == "✓")
        n_partial = sum(1 for r in results if r["judgment"] == "△")
        n_fail = sum(1 for r in results if r["judgment"] == "✗")
        print(f"  {name:<50} {len(results):>5} {n_pass:>5} {n_partial:>3} {n_fail:>3}")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Human-in-the-loop judge: load task rubrics, show deliverable content,
and guide through scoring each rubric item on the deliverable file.

Usage: uv run python scripts/judge_deliverables.py
       uv run python scripts/judge_deliverables.py --task sc_1ab7a50b23854581
"""

import sys
import json
import argparse
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

ACCEPTED_DIR = Path("data/accepted")
DELIVERABLES_DIR = Path("data/deliverables")


def read_docx_summary(path: Path) -> dict:
    from docx import Document
    doc = Document(str(path))
    paras = [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]

    text_content = "\n".join(t for _, t in paras)
    tables_info = []
    for i, t in enumerate(doc.tables):
        rows_data = []
        for ri, row in enumerate(t.rows):
            cells = [c.text.strip()[:60] for c in row.cells]
            rows_data.append(cells)
        tables_info.append({"rows": len(t.rows), "cols": len(t.columns), "data": rows_data[:8]})

    return {
        "paragraphs": len(paras),
        "total_chars": sum(len(t) for _, t in paras),
        "has_tables": len(doc.tables) > 0,
        "table_count": len(doc.tables),
        "tables": tables_info,
        "headings": [t for s, t in paras if s.startswith("Heading")],
        "full_text": text_content,
    }


def read_xlsx_summary(path: Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=False)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50)):
            row_data = []
            for c in row:
                if c.value is not None:
                    v = c.value
                    if isinstance(v, str) and v.startswith("="):
                        row_data.append({"ref": c.coordinate, "formula": v})
                    else:
                        row_data.append({"ref": c.coordinate, "value": str(v)[:60]})
            if row_data:
                rows.append(row_data)
        sheets.append({
            "name": name,
            "rows": ws.max_row,
            "cols": ws.max_column,
            "formulas": sum(1 for r in rows for c in r if "formula" in c),
            "data": rows[:15],
        })
    return {"sheets": sheets, "n_sheets": len(sheets)}


def read_md_summary(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    lines = text.split("\n")
    headings = [l.strip() for l in lines if l.strip().startswith("#")]
    return {
        "total_lines": len(lines),
        "total_chars": len(text),
        "headings": headings,
        "headings_count": len(headings),
        "full_text": text,
    }


def summarize(path: Path) -> dict:
    fmt = path.suffix.lower()
    summary = {"format": fmt, "size": path.stat().st_size, "path": str(path)}
    try:
        if fmt == ".docx":
            summary.update(read_docx_summary(path))
        elif fmt == ".xlsx":
            summary.update(read_xlsx_summary(path))
        elif fmt == ".md":
            summary.update(read_md_summary(path))
        else:
            summary["error"] = f"unsupported format: {fmt}"
    except Exception as e:
        summary["error"] = str(e)
    return summary


CATEGORY_NAMES = {
    "format": "格式/格式要求",
    "structure": "结构/章节要求",
    "style": "文风/措辞要求",
    "content": "内容要求",
    "accuracy": "准确性要求",
    "reference": "引用/参考文献要求",
    "data": "数据要求",
    "completeness": "完整性要求",
    "judgment": "判断/主观评分",
}


def judge_deliverable(task_id: str, task: dict, deliverable_path: Path):
    rubric = task.get("rubric", [])
    prompt = task.get("prompt", "")
    occupation = task.get("occupation", "?")
    archetype = task.get("archetype", "?")
    fmt = deliverable_path.suffix.lstrip(".")

    total_pts = sum(r["score"] for r in rubric)
    max_possible = sum(r["score"] for r in rubric if r["score"] > 0)
    max_penalty = sum(r["score"] for r in rubric if r["score"] < 0)

    # Read deliverable summary
    summary = summarize(deliverable_path)

    print(f"\n{'='*80}")
    print(f"  【{task_id}】{occupation} — {archetype} ({fmt.upper()})")
    print(f"  Deliverable: {deliverable_path.name} ({summary['size']:,} bytes)")
    print(f"  Rubric: {len(rubric)} items, total {total_pts}pts (max {max_possible}, penalty {max_penalty})")
    print(f"{'='*80}")

    # Print prompt (short)
    print(f"\n  ── Prompt (truncated) ──")
    print(f"  {prompt[:500]}")
    if len(prompt) > 500:
        print(f"  ... ({len(prompt)} chars total)")

    # Print deliverable content summary
    print(f"\n  ── Deliverable Content ──")
    if fmt == "docx":
        print(f"  Paragraphs: {summary.get('paragraphs', 0)}")
        print(f"  Total chars: {summary.get('total_chars', 0):,}")
        print(f"  Tables: {summary.get('table_count', 0)}")
        if summary.get("headings"):
            print(f"  Headings:")
            for h in summary["headings"]:
                print(f"    • {h}")
        if summary.get("tables"):
            for t in summary["tables"]:
                print(f"  Table: {t['rows']}x{t['cols']}")
                for row in t.get("data", [])[:4]:
                    print(f"    {' | '.join(str(c)[:40] for c in row)}")
    elif fmt == "xlsx":
        for s in summary.get("sheets", []):
            print(f"  Sheet: {s['name']} ({s['rows']}r x {s['cols']}c, {s['formulas']} formulas)")
    elif fmt == "md":
        print(f"  Lines: {summary.get('total_lines', 0)}")
        print(f"  Headings ({summary.get('headings_count', 0)}):")
        for h in summary.get("headings", []):
            print(f"    • {h}")

    # Judge each rubric item
    print(f"\n  ── Rubric Judgment ──")
    print(f"  ✓ = 满足 | ✗ = 不满足 | △ = 部分满足/需确认 | — = 无法从交付物判断")
    print(f"  {'':>4} {'Score':>5}  {'Cat':<12} {'Criterion'}")
    print(f"  {'':>4} {'-----':>5}  {'---':<12} {'---------'}")

    for i, r in enumerate(rubric):
        score = r["score"]
        cat = r.get("category", "?")
        criterion = r.get("criterion", "")
        is_penalty = r.get("is_penalty", False)
        score_str = f"{'+' if score >= 0 else ''}{score}"
        cat_display = f"{cat} {'[P]' if is_penalty else ''}"

        # Print item (judge is the human)
        print(f"  {'':>4} {score_str:>5}  {cat_display:<12} {criterion[:90]}")

    # Summary
    print(f"\n  ── Summary ──")
    print(f"  Rubric: {len(rubric)} items | Total: {total_pts} pts")
    print(f"  By category:")
    cat_counts = {}
    for r in rubric:
        cat = r.get("category", "?")
        cat_counts[cat] = cat_counts.get(cat, 0) + 1
    for cat, count in sorted(cat_counts.items(), key=lambda x: -x[1]):
        print(f"    {cat}: {count} items")

    # Input attachments
    print(f"\n  ── Input Attachments ──")
    input_paths = task.get("input_attachment_paths", []) or []
    if input_paths:
        for ip in input_paths:
            p = Path(ip)
            if p.exists():
                print(f"  ✅ {p.name} ({p.stat().st_size:,} bytes)")
            else:
                print(f"  ❌ {ip} (NOT FOUND)")
    else:
        print(f"  (none)")

    print()
    return summary


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--task", type=str, help="Specific task ID to judge")
    parser.add_argument("--occupation", type=str, help="Filter by occupation")
    args = parser.parse_args()

    task_files = sorted(ACCEPTED_DIR.glob("sc_*.json"))
    if not task_files:
        print("No accepted tasks found.")
        return

    # Allow filtering
    if args.task:
        task_files = [tf for tf in task_files if tf.stem == args.task]

    print(f"Found {len(task_files)} tasks to judge")

    all_summaries = {}
    for tf in task_files:
        task_id = tf.stem
        task = json.loads(tf.read_text(encoding="utf-8"))

        if args.occupation and task.get("occupation") != args.occupation:
            continue

        deliv_dir = DELIVERABLES_DIR / task_id
        if not deliv_dir.exists():
            print(f"\n⚠  {task_id}: No deliverable directory — SKIPPED")
            continue

        files = list(deliv_dir.iterdir())
        if not files:
            print(f"\n⚠  {task_id}: Empty deliverable directory — SKIPPED")
            continue

        for f in files:
            s = judge_deliverable(task_id, task, f)
            all_summaries[f"{task_id}/{f.name}"] = s

    # Quick overview table
    print(f"\n{'='*80}")
    print(f"  OVERVIEW")
    print(f"{'='*80}")
    print(f"  {'File':<50} {'Format':<6} {'Size':>10} {'Paras/Sheets/Lines':<20} {'Issues'}")
    print(f"  {'-'*90}")
    for name, s in sorted(all_summaries.items()):
        fmt = s.get("format", "?").lstrip(".")
        size = f"{s['size']:,}"
        info = ""
        if fmt == "docx":
            info = f"{s.get('paragraphs',0)}p/{s.get('table_count',0)}t"
        elif fmt == "xlsx":
            info = f"{s.get('n_sheets',0)}sh"
        elif fmt == "md":
            info = f"{s.get('total_lines',0)}l"
        error = " ⚠ ERROR" if s.get("error") else ""
        print(f"  {name:<50} {fmt:<6} {size:>10} {info:<20} {error}")


if __name__ == "__main__":
    main()
